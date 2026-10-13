"""文件说明：P2真实PG隔离schema最小测试；结束移除随机schema，样本保留var/tests供检查。"""

import os
from dataclasses import dataclass, field
from pathlib import Path
from uuid import uuid4
from unittest.mock import patch, Mock

import pytest
from dotenv import dotenv_values
from sqlalchemy import create_engine, text, select
from sqlalchemy.schema import CreateSchema, DropSchema
from knowledge.storage.schema import metadata
from knowledge.storage.repository import CoreRepository
from knowledge.config import KnowledgePaths
from knowledge.parser.database import ParserDatabase
from knowledge.parser.pipeline import ParserConfig, PolicyParserPipeline
from knowledge.parser.review import ReviewService
from knowledge.metadata_service import KnowledgeMetadataService
from knowledge.chunking.chunker import PolicyChunker
from knowledge.chunking.review import ChunkReviewService
from knowledge.indexing.database import IndexDatabase
from knowledge.indexing.config import IndexConfig, EmbeddingContract, canonical_json, sha256_text
from knowledge.catalog import KnowledgeCatalogService
from knowledge.read_service import KnowledgeReadService
from knowledge.indexing.indexer import PolicyIndexer
from knowledge.indexing.qdrant_store import StoredPoint
from knowledge.learning_source import freeze_learning_source

ROOT = Path(__file__).resolve().parents[2]


@dataclass
class TestDatabaseContext:
    """失败输出隐藏连接口令，仍支持测试中的解包。"""
    __test__ = False
    url: str = field(repr=False)
    engine: object = field(repr=False)
    paths: KnowledgePaths

    def __iter__(self):
        return iter((self.url, self.engine, self.paths))


@pytest.fixture
def isolated_pg(monkeypatch):
    values = dotenv_values(ROOT / "var/postgres-dev/platform.env")
    owner = create_engine(values["KNOWLEDGE_MIGRATION_URL"])
    schema = "p2_test_" + uuid4().hex
    monkeypatch.setenv("POLICY_INDEX_PROFILE", schema)
    monkeypatch.setenv("POLICY_INDEX_ALIAS", schema + "_active")
    monkeypatch.setenv("POLICY_COLLECTION_PREFIX", schema)
    with owner.begin() as connection:
        assert connection.execute(text("SELECT shobj_description(oid,'pg_database') FROM pg_database WHERE datname=current_database()" )).scalar() == "policy-platform-pg17-dev:knowledge"
        connection.execute(CreateSchema(schema))
        connection.execute(text(f'SET LOCAL search_path TO "{schema}"'))
        metadata.create_all(connection)
        connection.execute(text(f'GRANT USAGE ON SCHEMA "{schema}" TO knowledge_app'))
        connection.execute(text(f'GRANT SELECT,INSERT,UPDATE,DELETE ON ALL TABLES IN SCHEMA "{schema}" TO knowledge_app'))
        connection.execute(text(f'GRANT USAGE,SELECT ON ALL SEQUENCES IN SCHEMA "{schema}" TO knowledge_app'))
    engine = create_engine(values["KNOWLEDGE_DATABASE_URL"], connect_args={"options": f"-c search_path={schema}"})
    monkeypatch.setattr("knowledge.storage.repository.engine_for", lambda url: engine)
    monkeypatch.setenv("KNOWLEDGE_DATABASE_URL", values["KNOWLEDGE_DATABASE_URL"])
    root = ROOT / "var/tests/pg-p2"
    root.mkdir(parents=True, exist_ok=True)
    directory = root / uuid4().hex
    directory.mkdir()
    try:
        yield TestDatabaseContext(values["KNOWLEDGE_DATABASE_URL"], engine, KnowledgePaths(directory))
    finally:
        engine.dispose()
        with owner.begin() as connection:
            assert schema.startswith("p2_test_") and len(schema) == 40
            connection.execute(DropSchema(schema, cascade=True))
        owner.dispose()


def test_repository_short_reads_and_rollback(isolated_pg):
    url, engine, paths = isolated_pg
    with CoreRepository(url) as repo:
        assert repo.run(text("SELECT NULL AS nullable, 1 AS number")).mappings().one() == {"nullable": None, "number": 1}
        assert engine.pool.checkedout() == 0
        table = metadata.tables["policy_series"]
        value = repo.run(table.insert().values(series_key="rollback", canonical_title="中文", status="ACTIVE", created_at="2026-10-10T00:00:00+00:00", updated_at="2026-10-10T00:00:00+00:00").returning(table.c.id)).scalar_one()
        assert isinstance(value, int)
        repo.discard()
        assert repo.run(select(table)).all() == []
        from sqlalchemy.exc import IntegrityError
        with pytest.raises(IntegrityError):
            repo.run(table.insert().values(series_key="missing-required-fields"))
        assert engine.pool.checkedout() == 0
        assert repo.run(select(table)).all() == []


def test_actual_docx_parse_without_renderer(isolated_pg):
    from docx import Document
    url, engine, paths = isolated_pg
    paths.pending_documents.mkdir(parents=True)
    source = paths.pending_documents / "小样本.docx"
    doc = Document()
    doc.add_paragraph("测试制度")
    doc.add_paragraph("第一条 总则")
    doc.add_paragraph("本办法适用于本单位内部管理工作。")
    doc.save(source)
    with PolicyParserPipeline(url, paths.data_root, ParserConfig(render_pdf=False), progress=lambda _: None) as pipeline:
        result = pipeline.parse(source)
        assert result["status"] in {"succeeded", "review_required"}, result
    with engine.connect() as connection:
        assert connection.execute(text("SELECT count(*) FROM source_block")).scalar() > 0
        assert connection.execute(text("SELECT count(*) FROM policy_node")).scalar() > 0
    assert not list(paths.data_root.rglob("*.db"))


def test_metadata_chunk_index_and_read_contract(isolated_pg):
    url, engine, paths = isolated_pg
    with ParserDatabase(url) as db:
        document_id = db.get_or_create_document({"source_file": "pending/test.docx", "file_name": "测试制度.docx", "file_size": 100, "mime_type": "application/vnd.openxmlformats-officedocument.wordprocessingml.document", "checksum": "p2-test"})
        run_id = db.create_parse_run(document_id, "0.1.4", "p2-test")
        db.update_document(document_id, title="中文测试制度", latest_parse_run_id=run_id)
        db.update_run(run_id, status="succeeded")
        nodes = [{"temp_key": "n1", "parent_temp_key": None, "node_type": "article", "node_number": "第一条", "title": "奖励标准", "raw_text": "第一条 一等奖奖励不超过10000元。", "content": "一等奖奖励不超过10000元。", "normalized_content": "一等奖奖励不超过10000元。", "level": 1, "path": "第一条", "path_text": "第一条 奖励标准", "order_no": 1, "parser_confidence": 1.0, "decision_method": "rule", "rule_id": "test", "candidate_types": [], "review_status": "confirmed", "source_spans": []}]
        nodes[0]["source_block_key"] = "b1"
        block = metadata.tables["source_block"]
        block_id = db.repository.run(block.insert().values(parse_run_id=run_id, temp_key="b1", part_type="main", block_type="paragraph", order_no=1, raw_text=nodes[0]["raw_text"], clean_text=nodes[0]["raw_text"], properties_json="{}", source_locator_json="{}", normalization_actions_json="[]").returning(block.c.id)).scalar_one()
        db.repository.save()
        db.insert_nodes(document_id, run_id, nodes, {"b1": block_id})
    with ReviewService(url) as review:
        review.approve_run(run_id, "p2-reviewer")
    with KnowledgeMetadataService(url) as service:
        draft = service.ensure_draft(document_id, "p2-user", source_type="PARSE_EXTRACTED", initial_document_status="effective")
        revision_id = draft["id"]
        service.update_draft(document_id, revision_id, {"title": "中文测试制度", "issuing_org": "测试机构", "owner_department": "测试部门", "publish_date": "2026-10-10"})
        service.submit(document_id, revision_id)
        service.review(document_id, revision_id, True, "reviewer", "通过")
        assert service.list_series()[0]["document_count"] == 1
    with PolicyChunker(url) as chunker:
        chunk = chunker.run_document(document_id)
    assert chunk["status"] != "failed", chunk
    with ChunkReviewService(url) as review:
        with engine.connect() as connection:
            issues = list(connection.execute(text("SELECT id,severity FROM chunk_validation_issue WHERE resolved=0")).mappings())
        for issue in issues:
            assert issue["severity"] != "error", issue
            review.decide(chunk_run_id=chunk["id"], target_type="chunk_issue", target_id=issue["id"], action="confirm", reviewed_by="reviewer")
        review.approve_run(chunk["id"], "reviewer")
    config = IndexConfig.from_environment()
    with IndexDatabase(url) as index:
        profile = index.get_or_create_profile(config.profile_name, config.alias_name)
        sources = index.freeze_active_sources()
        assert sources[0]["document_id"] == document_id
        assert KnowledgeReadService(paths).review_detail("PARSE_RUN", run_id)
        assert KnowledgeReadService(paths).review_detail("CHUNK_RUN", chunk["id"])
        with pytest.raises(RuntimeError, match="active IndexRun"):
            KnowledgeCatalogService(paths).queryable_index_snapshot()
    # 仅替代外部模型/Qdrant；真实PG运行全部构建、核验和激活SQL。
    points = {}
    aliases = {}
    store = Mock()
    store.upsert.side_effect = lambda name, vector, batch: points.update({point.point_id: StoredPoint(point.point_id, point.vector, point.payload) for point in batch})
    store.exact_count.side_effect = lambda name: len(points)
    store.collection_exists.return_value = True
    store.retrieve.side_effect = lambda name, ids, **kw: {key: points[key] for key in ids}
    store.switch_alias.side_effect = lambda alias, name: aliases.update({alias: name})
    store.alias_target.side_effect = lambda alias: aliases.get(alias)
    embedding = Mock()
    embedding.embed.side_effect = lambda texts: [[1.0] * 8 for _ in texts]
    with PolicyIndexer(url, embedding_client=embedding, vector_store=store, contract=EmbeddingContract(dimension=8), config=config) as indexer:
        result = indexer.build(activate=False, required_document_ids=[document_id])
        # P3故障注入：Alias已切换、Knowledge提交失败，构建须冻结而非自动回切。
        with patch.object(indexer.db, "activate_run", side_effect=RuntimeError("模拟数据库提交失败")):
            with pytest.raises(RuntimeError, match="模拟数据库提交失败"):
                indexer.activate(result["index_run_id"])
        assert indexer.db.get_run(result["index_run_id"])["activation_status"] == "switching"
        with pytest.raises(RuntimeError, match="未收尾"):
            indexer.build(activate=False)
        switches = store.switch_alias.call_count
        target = aliases[config.alias_name]
        aliases[config.alias_name] = "different-target"
        with pytest.raises(ValueError, match="Alias尚未指向"):
            indexer.recover_activation(result["index_run_id"], base_run_id=None)
        assert indexer.summary(result["index_run_id"])["active_index_run_id"] is None
        aliases[config.alias_name] = target
        result = indexer.recover_activation(result["index_run_id"], base_run_id=None)
        indexer.recover_activation(result["index_run_id"], base_run_id=None)
        assert store.switch_alias.call_count == switches
        assert result["activation_status"] == "active"
        assert indexer.build(activate=False)["reused"]
        records = indexer.db.all_records(result["index_run_id"])
        ids = [row["chunk_id"] for row in records]
        assert indexer.db.hydrate_chunks(result["index_run_id"], ids)
        assert indexer.db.hydrate_citation_sources(result["index_run_id"], ids)
    catalog = KnowledgeCatalogService(paths)
    assert catalog.queryable_index_snapshot().documents[0].id == document_id
    assert catalog.get_document(document_id).status
    assert catalog.list_nodes(document_id)
    assert catalog.list_facets([document_id])["issuers"] == ["测试机构"]
    rows, total = catalog.list_documents(keyword="奖励", status=None, issuer=None, department=None, publish_date_from=None, publish_date_to=None, sort_by="updated_at", sort_order="desc", page=1, page_size=10, query_scope="CURRENT", allowed_document_ids=[document_id])
    assert total == 1 and rows[0].id == document_id
    assert KnowledgeReadService(paths).evidence_catalog("中文")
    assert freeze_learning_source(catalog, document_id, "FULL", [])
    # 新元数据草稿不得挤掉已发布快照，派生解析不能修改历史正文。
    from knowledge.parse_correction import ParseCorrectionService
    with KnowledgeMetadataService(url) as service:
        new_draft = service.ensure_draft(document_id, "reviewer", source_type="MANUAL_CORRECTION")
        assert new_draft["id"] != revision_id
    assert catalog.get_document(document_id).title == "中文测试制度"
    with engine.connect() as connection:
        node_id = connection.execute(text("SELECT id FROM policy_node ORDER BY id LIMIT 1")).scalar_one()
    with ParseCorrectionService(url) as correction:
        derived = correction.create(run_id, node_id, {"title": "修正标题"}, "reviewer")
        assert derived["parent_parse_run_id"] == run_id
    assert catalog.list_nodes(document_id)[1][0].title != "修正标题"


def test_api_startup_read_only():
    """真实新环境重复启动，只读查版本；不启动远程模型或后台轮询。"""
    from sqlalchemy import event
    from fastapi.testclient import TestClient
    from app.main import create_app
    from app.infrastructure.db.session import get_engine
    from knowledge.storage.repository import engine_for, database_url
    statements = []
    engines = [get_engine(), engine_for(database_url())]
    def capture(conn, cursor, statement, parameters, context, executemany):
        statements.append(statement.strip().split()[0].upper())
    for engine in engines:
        event.listen(engine, "before_cursor_execute", capture)
    try:
        for _ in range(2):
            with TestClient(create_app()) as client:
                assert client.get("/api/v1/health/live").status_code == 200
                assert client.get("/api/v1/health/ready").status_code == 200
        assert not set(statements) & {"CREATE", "ALTER", "DROP", "INSERT", "UPDATE", "DELETE", "TRUNCATE"}
    finally:
        for engine in engines:
            event.remove(engine, "before_cursor_execute", capture)


def test_production_paths_have_no_sqlite():
    import ast
    for root in (ROOT / "backend/knowledge", ROOT / "backend/app"):
        for path in root.rglob("*.py"):
            source = path.read_text(encoding="utf-8")
            tree = ast.parse(source)
            for node in ast.walk(tree):
                if isinstance(node, ast.Import):
                    assert all(alias.name != "sqlite3" for alias in node.names), path
            assert "POLICY_DB_PATH" not in source, path
            assert "lastrowid" not in source, path
            assert "BEGIN IMMEDIATE" not in source, path
