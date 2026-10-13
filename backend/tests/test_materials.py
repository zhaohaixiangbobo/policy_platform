"""文件说明：配套资料最小隔离回归；覆盖导入、版本、搜索、权限、CSRF和图片验证。"""
import io
import json
from dataclasses import replace
from pathlib import Path
from uuid import uuid4

import pytest
from fastapi import FastAPI
from fastapi.testclient import TestClient
from openpyxl import Workbook
from PIL import Image
from sqlalchemy import create_engine, select
from sqlalchemy.orm import sessionmaker
from sqlalchemy.pool import StaticPool

from app.api.dependencies import get_current_principal
from app.api.v1.materials import router
from app.application.auth_service import Principal, hash_token
from app.application.material_service import HEADERS
from app.config import Settings
from app.infrastructure.db.base import Base
from app.infrastructure.db.models import MaterialImportBatch, MaterialRevision, PlatformAuditLog
from app.infrastructure.db.session import get_db_session


def workbook(rows, headers=HEADERS):
    book = Workbook()
    sheet = book.active
    sheet.append(headers)
    for row in rows:
        sheet.append(row)
    output = io.BytesIO()
    book.save(output)
    return output.getvalue()


def row(text="原文：员工连续旷工五天，被解除劳动合同。", title="旷工处理案例"):
    return [title, "业务岗", "专管员", "一般人员", "", 44, "人事与劳动管理", "连续旷工", "解除劳动合同", text]


@pytest.fixture
def api(monkeypatch):
    root = Path(__file__).resolve(
    ).parents[2] / "var/tests" / ("materials-" + uuid4().hex)
    engine = create_engine(
        "sqlite://", connect_args={"check_same_thread": False}, poolclass=StaticPool)
    Base.metadata.create_all(engine)
    factory = sessionmaker(engine, expire_on_commit=False, autoflush=False)
    admin = Principal("admin", "admin", "管理员", ("ADMIN",), frozenset(
        {"material.read", "material.manage"}), False, "LOCAL", "session", hash_token("csrf"))
    identity = {"user": admin}
    app = FastAPI()
    app.include_router(router, prefix="/api/v1")
    app.dependency_overrides[get_current_principal] = lambda: identity["user"]

    def session():
        with factory() as value:
            yield value
    app.dependency_overrides[get_db_session] = session
    monkeypatch.setattr("app.api.v1.materials.get_settings",
                        lambda: Settings(materials_root=root))
    monkeypatch.setattr("app.api.dependencies.get_settings",
                        lambda: Settings(materials_root=root))
    with TestClient(app) as client:
        client.headers["X-CSRF-Token"] = "csrf"
        yield client, identity, factory, root
    engine.dispose()


BASE = "/api/v1/materials"


def test_v2_excel_normalization_and_merge():
    from app.application.material_service import parse_excel, MaterialError
    headers = [" 标题\n", "人员类别", "职务／岗位", *HEADERS[3:]]
    values = row()
    values[5] = 44.0
    values[8] = None
    result = parse_excel(workbook([values], headers))[0]
    assert result["error"] is None
    assert result["fields"]["age"] == "44"
    assert result["fields"]["outcome"] == ""
    assert result["fields"]["original"] == values[9]
    with pytest.raises(MaterialError, match="表头"):
        parse_excel(workbook([values[:7]], HEADERS[:7]))
    book = Workbook()
    book.active.append(HEADERS)
    book.active.append(values)
    book.active.merge_cells("B2:B3")
    stream = io.BytesIO()
    book.save(stream)
    with pytest.raises(MaterialError, match="B2:B3"):
        parse_excel(stream.getvalue())


def test_v2_schema_and_list_preview(api):
    from app.schemas.materials import CaseFields, ImageFields
    from pydantic import ValidationError
    with pytest.raises(ValidationError):
        ImageFields(title="图片", tags=["信访"], summary="说明", domains=["旧字段"])
    with pytest.raises(ValidationError):
        CaseFields(title="案例", original="案情", summary="旧字段")
    client, _, _, _ = api
    values = row(text="案情" * 100)
    values[8] = ""
    material_id = imported(client, [values])[0]
    item = client.get(BASE + "/" + material_id).json()
    assert item["fields"]["outcome"] == ""
    assert action(client, item).status_code == 200
    listed = client.get(BASE).json()["items"][0]
    assert listed["preview_text"] == values[9][:120]
    assert "original" not in listed["fields"]
    assert "outcome" not in listed["fields"]
    # 案情正文命中不能独立触发辅助推荐。
    assert client.get(BASE + "/recommendations",
                      params={"q": "案情"}).json()["items"] == []


def test_recommendations_published_version_and_offline(api):
    client, identity, _, _ = api
    material_id = imported(client, [row(title="信访举报流程")])[0]
    url = BASE + "/recommendations"
    assert client.get(url, params={"q": "信访举报怎么处理"}).json()["items"] == []
    item = client.get(BASE + "/" + material_id).json()
    assert action(client, item).status_code == 200
    result = client.get(url, params={"q": "请问信访举报如何办理"}).json()
    assert [i["id"] for i in result["items"]] == [material_id]
    published = result["items"][0]["revision_id"]
    item = client.get(BASE + "/" + material_id).json()
    fields = item["fields"] | {"title": "采购检查流程", "problem_types": ["采购检查"]}
    assert client.put(BASE + "/" + material_id, json={
                      "expected_version": item["version"], "fields": fields}).status_code == 200
    assert client.get(url, params={"q": "信访举报"}).json()[
        "items"][0]["revision_id"] == published
    assert client.get(url, params={"q": "采购检查"}).json()["items"] == []
    assert client.get(url, params={"q": "有什么流程图"}).json()["items"] == []
    item = client.get(BASE + "/" + material_id).json()
    assert action(client, item, "OFFLINE").status_code == 200
    assert client.get(url, params={"q": "信访举报"}).json()["items"] == []
    identity["user"] = replace(identity["user"], permissions=frozenset())
    assert client.get(url, params={"q": "信访举报"}).status_code == 403


def test_shared_material_search_rules():
    from app.application.material_search import SupportMaterialSearchService
    search = SupportMaterialSearchService()
    fields = {"title": "采购项目监督检查流程", "tags": ["采购", "监督检查", "流程"],
              "summary": "采购项目的检查步骤", "original": "这是一个偶然出现的劳动纪律词语。"}
    assert search.score("采购应该怎么监督检查", fields, "IMAGE",
                        recommend=True) >= search.threshold
    assert search.score("流程图", fields, "IMAGE", recommend=True) == 0
    assert search.score("流程", fields, "IMAGE") > 0
    assert search.score("劳动纪律", fields, "IMAGE") > 0
    assert search.score("劳动纪律", fields, "IMAGE", recommend=True) == 0
    assert search.score("信访举报流程", fields, "IMAGE", recommend=True) == 0
    # 领域仅为辅助信息，不能独立证明具体案例与问题相关。
    case = {"title": "某职工连续旷工问题", "domains": [
        "其他"], "problem_types": ["违反行规"], "original": "案情"}
    assert search.score("或者转属其他部门", case, "CASE", recommend=True) == 0
    assert search.score("物流运输方面的案例", case | {"domains": [
                        "物流运输"]}, "CASE", recommend=True) == 0


def preview(client, rows):
    return client.post(BASE + "/imports/preview", files={"file": ("案例.xlsx", workbook(rows))})


def imported(client, rows=None):
    response = preview(client, rows or [row()])
    assert response.status_code == 200, response.text
    result = client.post(
        BASE + f"/imports/{response.json()['batch_id']}/confirm")
    assert result.status_code == 200, result.text
    return result.json()["ids"]


def action(client, item, state="PUBLISH"):
    return client.post(BASE + "/actions", json={"action": state, "items": [{"id": item["id"], "expected_version": item["version"]}]})


def test_import_preserves_original_and_confirm_idempotent(api):
    client, _, factory, _ = api
    original = "  原始行\n保留换行和空格。  "
    p = preview(client, [row(original), row(original)]).json()
    assert (p["new"], p["duplicates"], p["errors"]) == (1, 1, 0)
    first = client.post(BASE + f"/imports/{p['batch_id']}/confirm").json()
    assert first["created"] == 1 and first["skipped"] == 1
    assert client.post(
        BASE + f"/imports/{p['batch_id']}/confirm").json() == first
    detail = client.get(BASE + "/" + first["ids"][0]).json()
    assert detail["fields"]["original"] == original
    assert detail["provenance"]["row"] == 2
    assert preview(client, [row(original)]).json()["duplicates"] == 1
    with factory() as session:
        assert len(session.scalars(select(PlatformAuditLog)).all()) >= 3


def test_validation_formula_blank_original_and_headers(api):
    client = api[0]
    p = preview(client, [row(""), row("=1+1")]).json()
    assert p["errors"] == 2
    assert client.post(
        BASE + f"/imports/{p['batch_id']}/confirm").status_code == 400
    bad = client.post(BASE + "/imports/preview",
                      files={"file": ("case.xlsx", workbook([row()], ["原文"]))})
    assert bad.status_code == 400


def test_version_isolation_publication_and_offline(api):
    client, identity, _, _ = api
    admin = identity["user"]
    item_id = imported(client)[0]
    draft = client.get(BASE + "/" + item_id).json()
    identity["user"] = replace(admin, permissions=frozenset({"material.read"}))
    assert client.get(BASE).json()["total"] == 0
    assert client.get(BASE + "/" + item_id).status_code == 404
    identity["user"] = admin
    assert action(client, draft).status_code == 200
    published = client.get(BASE + "/" + item_id).json()
    changed = {**published["fields"], "title": "新的标题"}
    saved = client.put(BASE + "/" + item_id,
                       json={"expected_version": published["version"], "fields": changed}).json()
    assert saved["revision_no"] == 2 and saved["state"] == "PUBLISHED"
    assert client.get(BASE + "?state=DRAFT").json()["total"] == 1
    identity["user"] = replace(admin, permissions=frozenset({"material.read"}))
    assert client.get(BASE + "/" + item_id).json()["fields"]["title"] != "新的标题"
    assert client.get(BASE + "/" + item_id,
                      params={"revision_id": saved["revision_id"]}).status_code == 404
    identity["user"] = admin
    assert action(client, saved).status_code == 200
    latest = client.get(BASE + "/" + item_id).json()
    old = client.get(BASE + "/" + item_id,
                     params={"revision_id": published["revision_id"]}).json()
    assert old["fields"]["title"] == published["fields"]["title"]
    assert old["revision_state"] == "SUPERSEDED"
    assert action(client, latest, "OFFLINE").status_code == 200
    identity["user"] = replace(admin, permissions=frozenset({"material.read"}))
    assert client.get(BASE + "/" + item_id).status_code == 404


def test_noop_save_stale_version_and_atomic_batch(api):
    client = api[0]
    ids = imported(client, [row(), row("第二条原文", "第二条")])
    one = client.get(BASE + "/" + ids[0]).json()
    noop = client.put(
        BASE + "/" + ids[0], json={"expected_version": one["version"], "fields": one["fields"]}).json()
    assert noop["version"] == one["version"] and noop["revision_no"] == 1
    assert client.put(
        BASE + "/" + ids[0], json={"expected_version": 999, "fields": one["fields"]}).status_code == 409
    assert client.put(BASE + "/" + ids[0], json={"expected_version": one["version"], "fields": {
                      **one["fields"], "original": "被篡改"}}).status_code == 400
    bad = client.post(BASE + "/actions", json={"action": "PUBLISH", "items": [
                      {"id": ids[0], "expected_version": 1}, {"id": ids[1], "expected_version": 999}]})
    assert bad.status_code == 409
    assert client.get(BASE + "/" + ids[0]).json()["state"] == "DRAFT"


def test_search_filter_pagination_and_publication_requirements(api):
    client = api[0]
    ids = imported(client, [row(f"唯一原文{i} 包含奖金", f"测试{i}") for i in range(11)])
    data = client.get(
        BASE, params={"state": "DRAFT", "q": "奖金", "domain": "人事与劳动管理"}).json()
    assert data["total"] == 11 and len(data["items"]) == 10
    assert "奖金" in data["items"][0]["match_excerpt"]
    assert len(client.get(BASE + "?state=DRAFT&page=2").json()["items"]) == 1
    assert client.get(BASE + "?state=DRAFT&q=%25").json()["total"] == 0


def test_permissions_csrf_and_batch_owner(api):
    client, identity, _, _ = api
    p = preview(client, [row()]).json()
    admin = identity["user"]
    identity["user"] = replace(admin, user_id="other-admin")
    assert client.post(
        BASE + f"/imports/{p['batch_id']}/confirm").status_code == 404
    identity["user"] = replace(admin, permissions=frozenset({"material.read"}))
    assert preview(client, [row()]).status_code == 403
    assert client.get(BASE + "?state=").status_code == 403
    identity["user"] = admin
    client.headers.pop("X-CSRF-Token")
    assert preview(client, [row()]).status_code == 403


def test_image_validation_dedup_and_file_access(api):
    client, identity, _, _ = api
    admin = identity["user"]
    stream = io.BytesIO()
    Image.new("RGB", (20, 40), "white").save(stream, format="PNG")
    response = client.post(
        BASE + "/images", files={"file": ("信访流程.png", stream.getvalue(), "image/png")}, data={"metadata": json.dumps({"title": "信访流程", "tags": ["信访举报"], "summary": "信访流程说明"})})
    assert response.status_code == 201, response.text
    item = response.json()
    assert item["fields"]["title"] == "信访流程"
    assert "key" not in item["file"]
    assert client.get(item["file_url"]).status_code == 200
    assert client.post(
        BASE + "/images", files={"file": ("同图.png", stream.getvalue())}, data={"metadata": json.dumps({"title": "同图", "tags": ["信访举报"], "summary": "说明"})}).status_code == 409
    identity["user"] = replace(admin, permissions=frozenset({"material.read"}))
    assert client.get(item["file_url"]).status_code == 404
    identity["user"] = admin
    assert client.post(
        BASE + "/images", files={"file": ("fake.png", b"<script>hi</script>")}).status_code == 400
    updated = client.put(BASE + "/" + item["id"], json={"expected_version": item["version"], "fields": {
                         **item["fields"], "tags": ["信访举报"], "summary": "信访流程说明"}}).json()
    assert action(client, updated).status_code == 200
    identity["user"] = replace(admin, permissions=frozenset({"material.read"}))
    file = client.get(item["file_url"])
    assert file.status_code == 200 and file.headers["x-content-type-options"] == "nosniff"


def test_expired_preview(api):
    from datetime import timedelta
    from app.infrastructure.db.models import utc_now
    client, _, factory, _ = api
    p = preview(client, [row()]).json()
    with factory() as session:
        session.get(MaterialImportBatch,
                    p["batch_id"]).expires_at = utc_now() - timedelta(seconds=1)
        session.commit()
    assert client.post(
        BASE + f"/imports/{p['batch_id']}/confirm").status_code == 409


def test_excel_limits_and_xml_entities(api):
    import zipfile
    client = api[0]
    assert preview(client, [row()] * 201).status_code == 400
    original = workbook([row()])
    target = io.BytesIO()
    with zipfile.ZipFile(io.BytesIO(original)) as source, zipfile.ZipFile(target, "w") as archive:
        for entry in source.infolist():
            data = source.read(entry.filename)
            if entry.filename == "xl/workbook.xml":
                data = b'<!DOCTYPE workbook [<!ENTITY x "unsafe">]>' + data
            archive.writestr(entry.filename, data)
    result = client.post(BASE + "/imports/preview",
                         files={"file": ("case.xlsx", target.getvalue())})
    assert result.status_code == 400 and "XML" in result.json()["detail"]


def test_keyword_ranking_and_unsupported_type(api):
    client = api[0]
    one = row("普通原文", "奖金标题")
    two = row("包含奖金的原文", "其他标题")
    imported(client, [one, two])
    result = client.get(
        BASE, params={"state": "DRAFT", "q": "奖金", "sort": "relevance"})
    assert result.json()["items"][0]["fields"]["title"] == "奖金标题"
    assert client.get(BASE + "?kind=TEMPLATE").status_code == 422
