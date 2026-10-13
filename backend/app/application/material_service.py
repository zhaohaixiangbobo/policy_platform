"""文件说明：配套资料版本、发布、关键词检索和Excel预览；不访问制度库、模型或Qdrant。"""
import hashlib
import io
import json
import re
import zipfile
import unicodedata
from datetime import date, datetime
from xml.etree import ElementTree
from datetime import timedelta, timezone
from uuid import uuid4

from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import select

from app.infrastructure.db.models import SupportingMaterial as Material, MaterialRevision as Revision, MaterialImportBatch as Batch, PlatformAuditLog, utc_now
from app.schemas.materials import CaseFields, ImageFields
from app.application.material_search import SupportMaterialSearchService, material_intent

HEADERS = ["标题", "人员类别", "职务/岗位", "职级/岗级", "政治面貌",
           "涉事年龄", "涉及领域", "问题性质", "处理结果", "案情描述"]
KEYS = ["title", "person_category", "position", "rank", "political_status",
        "age", "domains", "problem_types", "outcome", "original"]
# 字段中文名：案例沿用十列表头，图片补充标签和简介，用于把校验错误转成中文提示。
FIELD_LABELS = dict(zip(KEYS, HEADERS)) | {"tags": "标签/关键词", "summary": "简介"}


def describe_errors(exc):
    """把pydantic校验错误转换为“字段名+原因”的中文提示，不回显用户输入。"""
    parts = []
    for e in exc.errors():
        key = e["loc"][0] if e["loc"] else ""
        label = FIELD_LABELS.get(key, str(key) or "资料")
        kind = e["type"]
        if kind in {"string_too_short", "too_short", "missing"}:
            reason = "不能为空"
        elif kind in {"string_too_long", "too_long"}:
            reason = "超出长度或数量上限"
        elif kind == "extra_forbidden":
            reason = "不是当前资料类型允许的字段"
        elif kind == "value_error":
            # 自定义校验器已给出中文原因，去掉pydantic前缀即可。
            reason = "：" + str(e["msg"]).removeprefix("Value error, ")
        else:
            reason = "格式不正确"
        parts.append(label + reason)
    return "；".join(dict.fromkeys(parts))


def cell_text(value):
    if value is None:
        return ""
    # 纯日期单元格（零点）只保留日期，带时间的用空格分隔，避免出现“T00:00:00”。
    if isinstance(value, datetime):
        return value.date().isoformat() if value.time() == datetime.min.time() else value.isoformat(sep=" ")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def header_text(value):
    return re.sub(r"\s+", "", unicodedata.normalize("NFKC", cell_text(value)))


def encode(value):
    return json.dumps(value, ensure_ascii=False, separators=(",", ":"))


def fingerprint(kind, raw):
    return hashlib.sha256(kind.encode() + b":" + raw).hexdigest()


class MaterialError(ValueError):
    def __init__(self, message, status=400):
        super().__init__(message)
        self.status = status


def parse_excel(raw):
    """先限制ZIP展开量并拒绝XML实体和公式，再解析标准十列表；每批最多200行。"""
    if len(raw) > 10 * 1024 * 1024:
        raise MaterialError("Excel最多10MB")
    try:
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            entries = archive.infolist()
            if len(entries) > 1000 or sum(e.file_size for e in entries) > 30 * 1024 * 1024:
                raise MaterialError("Excel展开体积过大")
            for entry in entries:
                if entry.filename.endswith((".xml", ".rels")):
                    text = archive.read(entry).replace(b"\x00", b"").upper()
                    if b"<!DOCTYPE" in text or b"<!ENTITY" in text:
                        raise MaterialError("Excel包含不支持的XML声明")
                    if entry.filename.startswith("xl/worksheets/"):
                        root = ElementTree.fromstring(archive.read(entry))
                        merges = [e.attrib.get("ref", "") for e in root.iter(
                        ) if e.tag.endswith("}mergeCell")]
                        if merges:
                            raise MaterialError(
                                "请取消合并单元格并补齐内容：" + "、".join(merges))
        book = load_workbook(io.BytesIO(raw), read_only=True,
                             data_only=False, keep_links=False)
        try:
            sheet = book.worksheets[0]
            sheet.reset_dimensions()
            rows = sheet.iter_rows()
            head = next(rows, ())
            if [header_text(c.value) for c in head[:10]] != [header_text(h) for h in HEADERS] or any(c.value is not None for c in head[10:]):
                raise MaterialError("第一个工作表的表头必须依次为：" + "、".join(HEADERS))
            result = []
            for row_no, cells in enumerate(rows, 2):
                if row_no > 201:
                    raise MaterialError("每批最多200行，请拆分文件")
                if all(c.value is None for c in cells):
                    continue
                error = None
                fields = None
                if any(c.data_type in {"f", "e"} for c in cells):
                    error = "包含公式或Excel错误，请转换为纯文本"
                elif any(c.value is not None for c in cells[10:]):
                    error = "只接受标准十列"
                else:
                    values = {key: cell_text(cells[i].value) if i < len(
                        cells) else "" for i, key in enumerate(KEYS)}
                    original = values["original"]
                    for key in ("domains", "problem_types"):
                        values[key] = [v.strip() for v in re.split(
                            "[；;]", values[key]) if v.strip()]
                    try:
                        fields = CaseFields.model_validate(
                            values).model_dump()
                        fields["original"] = original  # 源文本逐字保存，去重也使用相同文本。
                    except ValidationError as exc:
                        # 案情为空白时由original的min_length拦截，提示为“案情描述不能为空”。
                        error = describe_errors(exc)
                result.append({"row": row_no, "fields": fields,
                              "error": error, "sheet": sheet.title})
            if not result:
                raise MaterialError("没有案例数据")
            return result
        finally:
            book.close()
    except MaterialError:
        raise
    except Exception as exc:
        raise MaterialError("无法读取Excel，请上传有效的.xlsx文件") from exc


class MaterialService:
    def __init__(self, session, principal, request_id="materials"):
        self.session = session
        self.principal = principal
        self.request_id = request_id
        self.manage = "material.manage" in principal.permissions

    def audit(self, action, object_id, details=None):
        self.session.add(PlatformAuditLog(request_id=self.request_id, actor_user_id=self.principal.user_id,
                                          auth_provider=self.principal.auth_provider, action="material." + action, object_type="supporting_material",
                                          object_id=object_id, status="SUCCEEDED", detail_json=encode(details or {})))

    def get(self, material_id):
        item = self.session.get(Material, material_id)
        if not item or (not self.manage and item.state != "PUBLISHED"):
            raise MaterialError("资料不存在或不可访问", 404)
        return item

    def revision(self, item, revision_id=None):
        target = revision_id or (
            item.latest_revision_id if self.manage else item.published_revision_id)
        if not self.manage and target != item.published_revision_id:
            raise MaterialError("资料版本不可访问", 404)
        rev = self.session.get(Revision, target)
        if not rev or rev.material_id != item.id:
            raise MaterialError("资料版本不存在", 404)
        return rev

    def public(self, item, rev, *, detail=False):
        payload = json.loads(rev.payload_json)
        file = payload.get("file")
        result = {"id": item.id, "kind": item.kind, "state": item.state, "version": item.version,
                  "revision_id": rev.id, "revision_no": rev.revision_no, "revision_state": rev.state,
                  "published_revision_id": item.published_revision_id, "has_draft": rev.state == "DRAFT",
                  "updated_at": item.updated_at.isoformat(), "published_at": rev.published_at.isoformat() if rev.published_at else None,
                  "fields": payload["fields"], "file_url": f"/api/v1/materials/{item.id}/revisions/{rev.id}/file" if file else None}
        if not detail:
            result["fields"] = {k: v for k, v in payload["fields"].items() if k not in {
                "original", "outcome"}}
        result["preview_text"] = (payload["fields"].get(
            "original", "") if item.kind == "CASE" else payload["fields"].get("summary", ""))[:120]
        if detail:
            result["provenance"] = payload.get("provenance", {})
            result["file"] = {k: v for k, v in (
                file or {}).items() if k != "key"}
        return result

    def fill_revision(self, rev, payload):
        f = payload["fields"]
        rev.title, rev.summary = f["title"], f.get("summary", "")
        rev.domains = ";" + ";".join(f.get("domains", [])) + ";"
        rev.keywords = ";".join(f.get("tags", []))
        rev.source_unit = rev.source_category = rev.region = rev.source_date = ""
        rev.search_text = "\n".join("；".join(v) if isinstance(
            v, list) else str(v) for k, v in f.items() if k != "visibility")
        rev.payload_json = encode(payload)

    def create(self, kind, fields, digest, *, file=None, provenance=None):
        if not self.manage:
            raise MaterialError("权限不足", 403)
        if kind not in {"CASE", "IMAGE"}:
            raise MaterialError("当前只支持案例和流程图/图片")
        if self.session.scalar(select(Material.id).where(Material.fingerprint == digest)):
            raise MaterialError("相同原始资料已经存在，请检索现有资料", 409)
        item = Material(id=str(uuid4()), kind=kind, fingerprint=digest,
                        latest_revision_id=str(uuid4()), created_by=self.principal.user_id)
        self.session.add(item)
        self.session.flush()
        rev = Revision(id=item.latest_revision_id, material_id=item.id,
                       revision_no=1, created_by=self.principal.user_id)
        self.fill_revision(
            rev, {"fields": fields, "file": file, "provenance": provenance or {}})
        self.session.add(rev)
        self.audit("create", item.id)
        self.session.flush()
        return item

    def detail(self, material_id, revision_id=None):
        item = self.get(material_id)
        rev = self.revision(item, revision_id)
        result = self.public(item, rev, detail=True)
        result["history"] = [{"id": r.id, "revision_no": r.revision_no, "state": r.state,
                              "created_at": r.created_at.isoformat()} for r in self.session.scalars(select(Revision).where(Revision.material_id == item.id).order_by(Revision.revision_no.desc()))] if self.manage else []
        return result

    def save(self, material_id, body):
        item = self.get(material_id)
        self.check_version(item, body.expected_version)
        rev = self.revision(item)
        payload = json.loads(rev.payload_json)
        updated = body.fields.model_dump()
        (CaseFields if item.kind == "CASE" else ImageFields).model_validate(updated)
        if item.kind == "CASE" and updated["original"].strip() != payload["fields"]["original"].strip():
            raise MaterialError("原始案例文本保留不变；不同原文请作为新资料导入")
        if item.kind == "CASE":
            updated["original"] = payload["fields"]["original"]
        if updated == payload["fields"]:
            return self.detail(item.id)
        payload["fields"] = updated
        if rev.state != "DRAFT":
            rev = Revision(id=str(uuid4()), material_id=item.id,
                           revision_no=rev.revision_no + 1, created_by=self.principal.user_id)
            self.session.add(rev)
            item.latest_revision_id = rev.id
        self.fill_revision(rev, payload)
        item.updated_at = utc_now()
        self.audit("save", item.id, {"revision_id": rev.id})
        self.session.flush()
        return self.detail(item.id)

    def check_version(self, item, expected):
        if not self.manage:
            raise MaterialError("权限不足", 403)
        if item.version != expected:
            raise MaterialError("资料已被修改，请刷新后重试", 409)

    def action(self, body):
        if len({i.id for i in body.items}) != len(body.items):
            raise MaterialError("批量操作包含重复资料")
        for selection in body.items:
            item = self.get(selection.id)
            self.check_version(item, selection.expected_version)
            rev = self.revision(item)
            if body.action == "PUBLISH":
                f = json.loads(rev.payload_json)["fields"]
                # 发布前按资料类型重新校验必填字段，提示中带上资料标题便于批量发布时定位。
                try:
                    (CaseFields if item.kind ==
                     "CASE" else ImageFields).model_validate(f)
                except ValidationError as exc:
                    raise MaterialError(
                        f"《{f.get('title') or item.id}》不满足发布要求：" + describe_errors(exc)) from exc
                if item.published_revision_id and item.published_revision_id != rev.id:
                    self.session.get(
                        Revision, item.published_revision_id).state = "SUPERSEDED"
                rev.state, rev.published_at = "PUBLISHED", utc_now()
                item.published_revision_id, item.state = rev.id, "PUBLISHED"
            else:
                if item.state != "PUBLISHED":
                    raise MaterialError("仅已发布资料可下架")
                item.state = "OFFLINE"
            item.updated_at = utc_now()
            self.audit(body.action.lower(), item.id, {"revision_id": rev.id})
        self.session.flush()
        return {"count": len(body.items)}

    def listing(self, *, q="", kind="", state="PUBLISHED", domain="", sort="updated", page=1):
        if not self.manage and state != "PUBLISHED":
            raise MaterialError("普通用户只能查看已发布资料", 403)
        # 管理员看最新编辑版本，普通读者始终看已发布版本。
        pointer = Material.latest_revision_id if self.manage else Material.published_revision_id
        conditions = ([Revision.state == "DRAFT"] if state == "DRAFT" else [
                      Material.state == state]) if state else []
        if domain:
            conditions.append(Revision.domains.contains(
                ";" + domain + ";", autoescape=True))
        tokens = q.split()[:10]
        base = select(Material, Revision).join(
            Revision, Revision.id == pointer).where(*conditions)
        order = [Revision.published_at.desc() if sort ==
                 "published" else Material.updated_at.desc(), Material.id]
        search = SupportMaterialSearchService()
        ranked = [(item, rev, search.score(q, json.loads(rev.payload_json)["fields"], item.kind))
                  for item, rev in self.session.execute(base.order_by(*order))]
        ranked = [r for r in ranked if not q.strip() or r[2] > 0]
        counts = {k: sum(r[0].kind == k for r in ranked)
                  for k in ("CASE", "IMAGE")}
        if kind:
            ranked = [r for r in ranked if r[0].kind == kind]
        if tokens and sort == "relevance":
            ranked.sort(key=lambda r: -r[2])
        total = len(ranked)
        rows = [(i, r) for i, r, _ in ranked[(page - 1) * 10:page * 10]]
        items = []
        for item, rev in rows:
            data = self.public(item, rev)
            pos = next((rev.search_text.lower().find(t.lower())
                       for t in tokens if t.lower() in rev.search_text.lower()), 0)
            data["match_excerpt"] = rev.search_text[max(
                0, pos - 30):pos + 140] if tokens else ""
            items.append(data)
        # 筛选选项同样只从当前用户可见版本生成。
        facet_query = select(Revision.domains).join(
            Material, pointer == Revision.id)
        if not self.manage:
            facet_query = facet_query.where(Material.state == "PUBLISHED")
        facets = self.session.execute(facet_query.distinct()).all()
        return {"items": items, "total": total, "page": page, "page_size": 10, "counts": counts,
                "domains": sorted({d for row in facets for d in row.domains.split(";") if d})}

    def recommend(self, query, kind=None):
        kind = kind or material_intent(query)
        # 与管理列表不同，即使管理员提问也仅查询已发布指针。
        search = SupportMaterialSearchService()
        ranked = []
        rows = self.session.execute(select(Material, Revision).join(
            Revision, Revision.id == Material.published_revision_id).where(
                Material.state == "PUBLISHED", Revision.state == "PUBLISHED"))
        for item, rev in rows:
            if kind and item.kind != kind:
                continue
            score = search.score(query, json.loads(rev.payload_json)[
                                 "fields"], item.kind, recommend=True)
            if score >= search.threshold:
                ranked.append((score, item, rev))
        ranked.sort(key=lambda r: (-r[0], r[1].id))
        items, counts = [], {"CASE": 0, "IMAGE": 0}
        for _, item, rev in ranked:
            if counts.get(item.kind, 2) < 2:
                result = self.public(item, rev)
                match = search.match(query, json.loads(rev.payload_json)["fields"], item.kind)
                result.update({key: match[key] for key in ("match_reason", "match_excerpt")})
                items.append(result)
                counts[item.kind] += 1
        return {"items": items, "search_version": search.version, "reference_only": True}

    def preview(self, raw, filename):
        rows = parse_excel(raw)
        seen = set()
        for row in rows:
            row["duplicate"] = False
            if row["error"]:
                continue
            digest = fingerprint(
                "CASE", row["fields"]["original"].encode("utf-8"))
            row["fingerprint"] = digest
            row["duplicate"] = digest in seen or bool(self.session.scalar(
                select(Material.id).where(Material.fingerprint == digest)))
            seen.add(digest)
        batch = Batch(created_by=self.principal.user_id, filename=filename, file_hash=hashlib.sha256(raw).hexdigest(),
                      payload_json=encode(rows), expires_at=utc_now() + timedelta(hours=24))
        self.session.add(batch)
        self.session.flush()
        self.audit("import_preview", batch.id)
        return {"batch_id": batch.id, "rows": rows, "total": len(rows), "errors": sum(bool(r["error"]) for r in rows),
                "duplicates": sum(r["duplicate"] for r in rows), "new": sum(not r["error"] and not r["duplicate"] for r in rows)}

    def confirm(self, batch_id):
        batch = self.session.get(Batch, batch_id)
        if not batch or batch.created_by != self.principal.user_id:
            raise MaterialError("导入预览不存在", 404)
        if batch.result_json:
            return json.loads(batch.result_json)
        if batch.expires_at.replace(tzinfo=timezone.utc) < utc_now():
            raise MaterialError("预览已过期，请重新上传", 409)
        rows = json.loads(batch.payload_json)
        if any(r["error"] for r in rows):
            raise MaterialError("请修正Excel中的错误行后重新预览")
        ids, skipped = [], 0
        for row in rows:
            if row["duplicate"] or self.session.scalar(select(Material.id).where(Material.fingerprint == row["fingerprint"])):
                skipped += 1
                continue
            item = self.create("CASE", row["fields"], row["fingerprint"], provenance={"batch_id": batch.id,
                                                                                      "filename": batch.filename, "file_hash": batch.file_hash, "sheet": row["sheet"], "row": row["row"]})
            ids.append(item.id)
        result = {"created": len(ids), "skipped": skipped, "ids": ids}
        batch.result_json = encode(result)
        self.audit("import_confirm", batch.id, result)
        self.session.flush()
        return result
